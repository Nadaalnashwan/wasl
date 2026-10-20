"""Upload validation, text extraction and storage (local disk or Supabase Storage)."""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path

import httpx
from docx import Document
from pypdf import PdfReader

from app.config import settings

ALLOWED = {
    ".pdf": ("application/pdf", b"%PDF"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK\x03\x04"),
    ".txt": ("text/plain", None),
    ".md": ("text/markdown", None),
}


class UploadRejected(ValueError):
    pass


@dataclass
class ExtractedDoc:
    filename: str
    ext: str
    mime: str
    size: int
    sha256: str
    text: str


def _safe_name(name: str) -> str:
    base = Path(name or "file").name
    return re.sub(r"[^\w.\- ]", "_", base)[:120] or "file"


def validate_and_extract(filename: str, data: bytes) -> ExtractedDoc:
    name = _safe_name(filename)
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED:
        raise UploadRejected(f"Unsupported file type '{ext or 'none'}'. Allowed: PDF, DOCX, TXT, MD.")
    limit = settings.max_upload_mb * 1024 * 1024
    if len(data) == 0:
        raise UploadRejected("The file is empty.")
    if len(data) > limit:
        raise UploadRejected(f"File is larger than {settings.max_upload_mb} MB.")
    mime, magic = ALLOWED[ext]
    if magic and not data.startswith(magic):
        raise UploadRejected("File content does not match its extension.")
    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise UploadRejected("Password-protected PDFs are not supported.")
            if len(reader.pages) > 30:
                raise UploadRejected("PDF has more than 30 pages.")
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif ext == ".docx":
            doc = Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(lines)
        else:
            text = data.decode("utf-8", errors="replace")
    except UploadRejected:
        raise
    except Exception as exc:
        raise UploadRejected("The document could not be read. Is it corrupted?") from exc
    text = text.replace("\x00", "").strip()
    if len(text) < 30:
        raise UploadRejected("No readable text found (scanned images are not supported yet — please upload a text-based PDF or DOCX).")
    return ExtractedDoc(name, ext, mime, len(data), hashlib.sha256(data).hexdigest(), text)


def store(user_id: str, doc: ExtractedDoc, data: bytes) -> str:
    """Persist the original file. Returns a storage path (never a public URL)."""
    key = f"{user_id}/{doc.sha256[:16]}{doc.ext}"
    if settings.supabase_url and settings.supabase_service_role_key:
        url = f"{settings.supabase_url.rstrip('/')}/storage/v1/object/{settings.supabase_storage_bucket}/{key}"
        resp = httpx.post(url, content=data, timeout=30, headers={
            "Authorization": f"Bearer {settings.supabase_service_role_key}", "Content-Type": doc.mime, "x-upsert": "true",
        })
        resp.raise_for_status()
        return f"supabase://{settings.supabase_storage_bucket}/{key}"
    path = Path(settings.upload_dir) / key
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return f"local://{key}"

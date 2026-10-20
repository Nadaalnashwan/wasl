"""CV export to DOCX (python-docx) and PDF (fpdf2 with HarfBuzz shaping for Arabic)."""
from __future__ import annotations

import io
import logging
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fpdf import FPDF

from app.ai.schemas import CVData

BRAND = (11, 79, 179)
_AR = re.compile(r"[\u0600-\u06FF]")

FONT_CANDIDATES = [
    os.environ.get("PDF_FONT_PATH", ""),
    "C:/Windows/Fonts/arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
]
BOLD_CANDIDATES = [
    os.environ.get("PDF_FONT_BOLD_PATH", ""),
    "C:/Windows/Fonts/arialbd.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]


def _first(paths: list[str]) -> str | None:
    return next((p for p in paths if p and Path(p).exists()), None)


def is_rtl(cv: CVData) -> bool:
    sample = " ".join([cv.contact.full_name, cv.summary, *(e.title for e in cv.experience)])
    return len(_AR.findall(sample)) > len(sample) * 0.2


def _sections(cv: CVData) -> list[tuple[str, list[tuple[str, str, list[str]]]]]:
    """Normalised (heading, [(title, meta, bullets)]) list shared by both exporters."""
    ar = is_rtl(cv)
    H = (lambda en, a: a) if ar else (lambda en, a: en)
    out = []
    if cv.experience:
        out.append((H("Experience", "الخبرات"), [
            (f"{e.title} — {e.organization}".strip(" —"), " · ".join(filter(None, [e.location, f"{e.start} – {e.end}".strip(" –")])), e.bullets)
            for e in cv.experience]))
    if cv.projects:
        out.append((H("Projects", "المشاريع"), [
            (p.name, ", ".join(p.technologies), ([p.description] if p.description else []) + p.bullets) for p in cv.projects]))
    if cv.education:
        out.append((H("Education", "التعليم"), [
            (f"{e.degree} — {e.institution}".strip(" —"), f"{e.start} – {e.end}".strip(" –"), [e.details] if e.details else [])
            for e in cv.education]))
    if cv.certifications:
        out.append((H("Certifications", "الشهادات"), [(c.name, " · ".join(filter(None, [c.issuer, c.date])), []) for c in cv.certifications]))
    if cv.awards:
        out.append((H("Awards", "الجوائز"), [(a.title, " · ".join(filter(None, [a.issuer, a.date])), [a.description] if a.description else []) for a in cv.awards]))
    return out


def to_docx(cv_dict: dict) -> bytes:
    cv = CVData.model_validate(cv_dict)
    rtl = is_rtl(cv)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    def para(text: str, size: float = 10.5, bold: bool = False, color: tuple | None = None, align=None):
        p = doc.add_paragraph()
        if rtl:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p._p.get_or_add_pPr().append(p._p.makeelement(qn("w:bidi"), {}))
        elif align is not None:
            p.paragraph_format.alignment = align
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        if rtl:
            run._r.get_or_add_rPr().append(run._r.makeelement(qn("w:rtl"), {}))
        if color:
            run.font.color.rgb = RGBColor(*color)
        p.paragraph_format.space_after = Pt(2)
        return p

    para(cv.contact.full_name or "Curriculum Vitae", 20, True, BRAND)
    if cv.contact.headline:
        para(cv.contact.headline, 12)
    contact = " · ".join(filter(None, [cv.contact.email, cv.contact.phone, cv.contact.location, *cv.contact.links]))
    if contact:
        para(contact, 9, color=(90, 90, 90))
    if cv.summary:
        para("الملخص" if rtl else "Summary", 13, True, BRAND)
        para(cv.summary)
    for heading, items in _sections(cv):
        para(heading, 13, True, BRAND)
        for title, meta, bullets in items:
            para(title, 11, True)
            if meta:
                para(meta, 9, color=(100, 100, 100))
            for b in bullets:
                para(f"• {b}")
    if cv.skills:
        para("المهارات" if rtl else "Skills", 13, True, BRAND)
        para(", ".join(cv.skills))
    if cv.languages:
        para("اللغات" if rtl else "Languages", 13, True, BRAND)
        para(", ".join(cv.languages))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(cv_dict: dict) -> bytes:
    cv = CVData.model_validate(cv_dict)
    rtl = is_rtl(cv)
    pdf = FPDF(format="A4")
    pdf.set_margins(16, 14, 16)
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()
    regular, bold = _first(FONT_CANDIDATES), _first(BOLD_CANDIDATES)
    if regular:
        pdf.add_font("Body", "", regular)
        pdf.add_font("Body", "B", bold or regular)
        family = "Body"
        try:
            pdf.set_text_shaping(True)
        except Exception:  # uharfbuzz missing -> Latin still fine, Arabic letters render unshaped
            logging.getLogger("wasl.export").warning("text shaping unavailable; install uharfbuzz for Arabic PDFs")
    else:
        family = "Helvetica"
    align = "R" if rtl else "L"
    width = pdf.epw

    def safe(t: str) -> str:
        return t if regular else t.encode("latin-1", "replace").decode("latin-1")

    def line(text: str, size: float = 10, style: str = "", color=(30, 30, 30), gap: float = 1.5):
        pdf.set_font(family, style, size)
        pdf.set_text_color(*color)
        pdf.multi_cell(width, size * 0.5, safe(text), align=align, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    def heading(text: str):
        pdf.ln(2)
        line(text.upper() if not rtl else text, 11.5, "B", BRAND, 0.5)
        pdf.set_draw_color(*BRAND)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + width, pdf.get_y())
        pdf.ln(2)

    line(cv.contact.full_name or "Curriculum Vitae", 20, "B", BRAND, 1)
    if cv.contact.headline:
        line(cv.contact.headline, 11.5)
    contact = " · ".join(filter(None, [cv.contact.email, cv.contact.phone, cv.contact.location, *cv.contact.links]))
    if contact:
        line(contact, 8.5, color=(100, 100, 100))
    if cv.summary:
        heading("الملخص" if rtl else "Summary")
        line(cv.summary)
    for h, items in _sections(cv):
        heading(h)
        for title, meta, bullets in items:
            line(title, 10.5, "B", gap=0.5)
            if meta:
                line(meta, 8.5, color=(100, 100, 100), gap=0.5)
            for b in bullets:
                line(f"• {b}", 9.5, gap=0.5)
            pdf.ln(1.5)
    if cv.skills:
        heading("المهارات" if rtl else "Skills")
        line(", ".join(cv.skills))
    if cv.languages:
        heading("اللغات" if rtl else "Languages")
        line(", ".join(cv.languages))
    return bytes(pdf.output())
